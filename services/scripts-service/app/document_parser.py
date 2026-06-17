"""Извлечение текста из документов скриптов: DOCX, PDF, RTF, TXT.

Используется при импорте полнотекстового скрипта. Возвращает плоский текст
с сохранением переносов строк — структуру по смыслу дальше строит LLM.
"""
from __future__ import annotations

import io
import logging

logger = logging.getLogger(__name__)

# Документ больше этого размера обрезается перед отправкой в LLM
MAX_TEXT_CHARS = 40_000

SUPPORTED_EXTENSIONS = (".docx", ".pdf", ".rtf", ".txt", ".md")


class DocumentParseError(Exception):
    pass


def parse_document(file_bytes: bytes, filename: str) -> str:
    """Определяет формат по расширению и извлекает текст."""
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError(
            f"Неподдерживаемый формат «{ext or filename}». Поддерживаются: DOCX, PDF, RTF, TXT, MD"
        )
    if not file_bytes:
        raise DocumentParseError("Пустой файл")

    if ext == ".docx":
        text = _parse_docx(file_bytes)
    elif ext == ".pdf":
        text = _parse_pdf(file_bytes)
    elif ext == ".rtf":
        text = _parse_rtf(file_bytes)
    else:
        text = _parse_txt(file_bytes)

    text = _normalize(text)
    if not text.strip():
        raise DocumentParseError("Не удалось извлечь текст из документа")
    if len(text) > MAX_TEXT_CHARS:
        logger.warning("Document %s truncated: %d > %d chars", filename, len(text), MAX_TEXT_CHARS)
        text = text[:MAX_TEXT_CHARS]
    return text


def _parse_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise DocumentParseError("python-docx не установлен") from e
    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise DocumentParseError(f"Не удалось открыть DOCX: {e}") from e

    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    # Таблицы тоже содержат текст скриптов (часто скрипты оформляют таблицей)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _parse_pdf(file_bytes: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise DocumentParseError("pypdf не установлен") from e
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as e:
        raise DocumentParseError(f"Не удалось прочитать PDF: {e}") from e


def _parse_rtf(file_bytes: bytes) -> str:
    try:
        from striprtf.striprtf import rtf_to_text
    except ImportError as e:
        raise DocumentParseError("striprtf не установлен") from e
    try:
        return rtf_to_text(file_bytes.decode("utf-8", errors="ignore"))
    except Exception as e:
        raise DocumentParseError(f"Не удалось прочитать RTF: {e}") from e


def _parse_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "cp1251"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def _normalize(text: str) -> str:
    lines = [ln.rstrip() for ln in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    out: list[str] = []
    blank = 0
    for ln in lines:
        if not ln.strip():
            blank += 1
            if blank > 1:
                continue
        else:
            blank = 0
        out.append(ln)
    return "\n".join(out).strip()
